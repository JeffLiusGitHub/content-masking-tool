"""Regression contracts for persistent rules, URL masking and review selections."""
import pytest

from maskingtool.denylist import loader
from maskingtool.engine import MaskingEngine
from maskingtool.operators import restore_text
from maskingtool.vault import Vault


@pytest.fixture
def rules(tmp_path, monkeypatch):
    monkeypatch.setenv("MASKINGTOOL_DATA_DIR", str(tmp_path / "app"))
    return tmp_path / "app" / "denylists"


def test_multiline_rule_preserves_exact_bytes_and_is_loaded(rules):
    term = " https://example.test/a\r\n下一行 "
    loader.add_deny_term(term, "URL")
    assert term in loader.load_deny_lists()["URL"]
    assert loader.load_manual_terms() == [{"term": term, "entity_type": "URL"}]
    assert not (rules / "urls.csv").exists()


def test_strict_undo_and_remask(rules):
    term = "https://example.test/private"
    loader.add_deny_term(term, "URL")
    with loader.rule_transaction():
        assert loader.undo_terms([term]) == {term}
    assert term not in loader.load_deny_lists().get("URL", [])
    assert loader.load_allow_terms() == []
    with loader.rule_transaction():
        assert loader.undo_terms([term]) == set()
    assert loader.load_allow_terms() == [term]
    loader.add_deny_term(term, "URL")
    assert loader.load_allow_terms() == []


def test_remove_cleans_csv_and_manual_duplicates_but_not_full_name(rules):
    loader.add_deny_term("John Smith", "PERSON")
    with loader.rule_transaction():
        loader.undo_terms(["John"])
    assert "John Smith" in loader.load_deny_lists()["PERSON"]
    assert loader.load_allow_terms() == ["John"]
    with loader.rule_transaction():
        loader.undo_terms(["John Smith"])
    assert loader.load_deny_lists()["PERSON"] == []
    assert loader.load_manual_terms() == []


def test_corrupt_allowlist_fails_closed(rules):
    rules.mkdir(parents=True)
    (rules / "allow_terms.json").write_text('{"terms": 12}', encoding="utf-8")
    with pytest.raises(ValueError):
        loader.load_allow_terms()


def test_rule_transaction_rolls_back_when_preview_fails(rules):
    loader.add_deny_term("Alice", "PERSON")
    before = {p.name: p.read_bytes() for p in rules.iterdir() if p.suffix != ".lock"}
    with pytest.raises(RuntimeError):
        with loader.rule_transaction():
            loader.undo_terms(["Alice", "Other"])
            raise RuntimeError("preview failed")
    after = {p.name: p.read_bytes() for p in rules.iterdir() if p.suffix != ".lock"}
    assert after == before


@pytest.mark.parametrize("url", [
    "https://example.test/a?x=1&y=two#part", "http://example.test/a%20b",
    "www.example.test/path", "https://example.test/a_(b)",
])
def test_url_priority_persistence_and_punctuation(url, tmp_path):
    v = Vault.create("urls", vaults_dir=tmp_path)
    text = f"({url}). Again {url}!"
    e = MaskingEngine({"ORG": ["example"]}, enable_ner=False)
    masked = e.mask_text(text, v)
    assert masked == "(⟦URL_001⟧). Again ⟦URL_001⟧!"
    v.save()
    assert restore_text(masked, Vault.load(v.vault_id, tmp_path)) == (text, [])


def test_url_no_bare_domains_no_cross_line_guessing(tmp_path):
    e = MaskingEngine({})
    v = Vault.create("urls", vaults_dir=tmp_path)
    assert e.mask_text("example.test", v) == "example.test"
    text = "https://example.test/\ncontinuation"
    masked = e.mask_text(text, v)
    assert masked.endswith("\ncontinuation")
    assert restore_text(masked, v) == (text, [])


def test_allowlist_only_protects_its_range(tmp_path):
    v = Vault.create("urls", vaults_dir=tmp_path)
    e = MaskingEngine({}, allow_terms=["example.test"])
    masked = e.mask_text("https://example.test/private", v)
    assert masked == "⟦URL_001⟧example.test⟦URL_002⟧"
    assert restore_text(masked, v)[0] == "https://example.test/private"


def test_manual_block_wins_and_preserves_crlf(tmp_path):
    term = "https://example.test/\r\n秘密"
    e = MaskingEngine({}, manual_terms=[{"term": term, "entity_type": "URL"}])
    v = Vault.create("multi", vaults_dir=tmp_path)
    assert e.mask_text(term, v) == "⟦URL_001⟧"
    assert restore_text("⟦URL_001⟧", v) == (term, [])


def test_result_selection_maps_tokens_and_non_bmp_characters(tmp_path):
    from maskingtool.selection import map_selection
    v = Vault.create("selection", vaults_dir=tmp_path)
    token = v.get_or_create_token("https://example.test", "ORG")
    before = "😀 " + "https://example.test/path\r\nnext"
    after = "😀 " + token + "/path\r\nnext"
    selection = map_selection(before, after, v, 4, len(after), masked=True)
    assert selection.text == "https://example.test/path\r\nnext"
    assert selection.masked_originals == ["https://example.test"]


def test_selection_batch_undo_targets_individual_originals(tmp_path):
    from maskingtool.selection import map_selection
    v = Vault.create("selection", vaults_dir=tmp_path)
    after = MaskingEngine({"PERSON": ["Alice", "Bob"]}).mask_text("Alice and Bob", v)
    selected = map_selection("Alice and Bob", after, v, 0, len(after), masked=True)
    assert selected.masked_originals == ["Alice", "Bob"]


def test_preview_undo_remask_and_confirm_preserves_old_output(rules, tmp_path):
    from maskingtool import config
    from maskingtool.gui_core import prepare_mask_preview, edit_preview, commit_mask_preview
    config.save_settings({"enable_ner": False})
    url = "https://example.test/a"
    p = tmp_path / "source.md"
    p.write_text(url + "\n" + url, encoding="utf-8")
    loader.add_deny_term(url, "URL")
    initial = prepare_mask_preview(p)
    old = commit_mask_preview(initial)
    old_bytes = old.output_path.read_bytes()
    undone = edit_preview(initial, undo=[url])
    assert undone.masked_text == undone.before_text
    assert loader.load_allow_terms() == []
    assert "⟦URL_001⟧" in prepare_mask_preview(p).masked_text
    remasked = edit_preview(undone, term=url, entity_type="URL")
    assert remasked.masked_text == initial.masked_text
    assert remasked.vault.vault_id == initial.vault.vault_id
    committed = commit_mask_preview(remasked)
    assert committed.output_path != old.output_path
    assert old.output_path.read_bytes() == old_bytes
    assert committed.entity_counts == {"URL": 1}
    assert commit_mask_preview(undone).entity_counts == {}


def test_automatic_unmask_is_persistent_and_counts_only_active(rules, tmp_path):
    from maskingtool import config
    from maskingtool.gui_core import prepare_mask_preview, edit_preview, commit_mask_preview
    config.save_settings({"enable_ner": False})
    p = tmp_path / "source.md"
    p.write_text("https://example.test/a https://example.test/b", encoding="utf-8")
    initial = prepare_mask_preview(p)
    result = edit_preview(initial, undo=["https://example.test/a"])
    assert result.masked_text == "https://example.test/a ⟦URL_002⟧"
    assert "https://example.test/a" in prepare_mask_preview(p).masked_text
    assert commit_mask_preview(result).entity_counts == {"URL": 1}


def test_write_failure_rolls_back_all_files(rules, monkeypatch):
    loader.add_deny_term("Alice", "PERSON")
    before = {p.name: p.read_bytes() for p in rules.iterdir() if p.suffix != ".lock"}
    replace = loader.os.replace
    failed = False
    def fail_once(src, dest):
        nonlocal failed
        if str(dest).endswith("allow_terms.json") and not failed:
            failed = True
            raise OSError("disk write failed")
        return replace(src, dest)
    monkeypatch.setattr(loader.os, "replace", fail_once)
    with pytest.raises(OSError):
        loader.undo_terms(["Alice", "automatic"])
    assert {p.name: p.read_bytes() for p in rules.iterdir() if p.suffix != ".lock"} == before


def test_failed_preview_does_not_change_rules_or_vault(rules, tmp_path, monkeypatch):
    from maskingtool import config, gui_core
    config.save_settings({"enable_ner": False})
    p = tmp_path / "source.md"
    p.write_text("Alice", encoding="utf-8")
    initial = gui_core.prepare_mask_preview(p)
    def fail(*a, **kw):
        raise RuntimeError("analysis failed")
    monkeypatch.setattr(gui_core, "prepare_mask_preview", fail)
    with pytest.raises(RuntimeError):
        gui_core.edit_preview(initial, term="Alice", entity_type="PERSON")
    assert "Alice" not in loader.load_deny_lists()["PERSON"]
    assert len(initial.vault) == 0


def test_manual_multiline_is_not_trimmed_by_fallback_recognizers(tmp_path):
    v = Vault.create("exact", vaults_dir=tmp_path)
    term = " first\r\nlast "
    e = MaskingEngine({"TEXT": [term]}, manual_terms=[{"term": term, "entity_type": "TEXT"}])
    assert e.mask_text("first\nlast", v) == "first\nlast"
    assert e.mask_text("first\r\nlast", v) == "first\r\nlast"
    assert e.mask_text(term, v) == "⟦TEXT_001⟧"


def test_allow_rule_is_not_a_substring_bypass(tmp_path):
    v = Vault.create("allow", vaults_dir=tmp_path)
    e = MaskingEngine({"PERSON": ["Alice", "Aliceland"]}, allow_terms=["Alice"])
    assert e.mask_text("Alice Aliceland", v) == "Alice ⟦PERSON_001⟧"


def test_manual_domain_does_not_expose_the_rest_of_a_url(tmp_path):
    v = Vault.create("url", vaults_dir=tmp_path)
    e = MaskingEngine({}, manual_terms=[{"term": "example.test", "entity_type": "ORG"}])
    source = "https://example.test/private?id=123"
    masked = e.mask_text(source, v)
    assert "/private" not in masked and "https://" not in masked
    assert "⟦ORG_001⟧" in masked
    assert restore_text(masked, v) == (source, [])


def test_allow_exact_selection_including_spaces_between_words(tmp_path):
    v = Vault.create("spaces", vaults_dir=tmp_path)
    term = " https://example.test/path "
    source = "before" + term + "after"
    assert MaskingEngine({}, allow_terms=[term]).mask_text(source, v) == source


def test_canonical_docx_table_does_not_insert_spaces_between_runs(tmp_path):
    import docx
    from maskingtool.pipeline import mask_file, read_document
    document = docx.Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].add_run("Acme ")
    cell.paragraphs[0].add_run("Research").bold = True
    cell.add_paragraph("Next paragraph")
    path = tmp_path / "runs.docx"
    document.save(path)
    source, _ = read_document(path)
    assert "Acme Research" in source
    v = Vault.create(path.name, vaults_dir=tmp_path / "vaults")
    result = mask_file(path, MaskingEngine({"ORG": ["Acme Research"]}), v)
    assert "Acme" not in result.masked_text and "Research" not in result.masked_text
    assert restore_text(result.masked_text, v) == (source, [])
