"""Milestone 5 tests: DOCX pipeline — extraction, split runs, formatting,
masked-docx write-back, restore. TESTPLAN.md section 4.5.

The fixture .docx is generated programmatically (no binary blob in the repo)
so its structure — including a name deliberately split across two runs — is
readable right here.
"""
from pathlib import Path

import docx
import pytest

from maskingtool.engine import MaskingEngine
from maskingtool.parsers.docx_parser import (
    mask_docx_to_markdown,
    parse_docx,
    write_masked_docx,
)
from maskingtool.renderers.docx_renderer import restore_docx
from maskingtool.vault import TOKEN_PATTERN, Vault

DENY = {
    "ORG": ["Acme Corp", "Bidco Ltd"],
    "PERSON": ["John Smith", "张三"],
}
ALL_TERMS = [t for terms in DENY.values() for t in terms]


def _doc_text(path: Path) -> str:
    """All text content of a docx: body paragraphs + table cells."""
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


@pytest.fixture
def fixture_docx(tmp_path) -> Path:
    doc = docx.Document()
    doc.add_heading("Acme Corp Review", level=1)

    p1 = doc.add_paragraph()
    p1.add_run("Prepared by ")
    bold = p1.add_run("John Smith")
    bold.bold = True
    p1.add_run(" for the board.")

    # name deliberately split across two runs (spellcheck artifact simulation)
    p2 = doc.add_paragraph()
    p2.add_run("Our client Acme ")
    p2.add_run("Corp expands, working with 张三.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Company"
    table.cell(0, 1).text = "Contact"
    table.cell(1, 0).text = "Bidco Ltd"
    table.cell(1, 1).text = "John Smith"

    doc.add_paragraph("Closing line with nothing sensitive.")
    path = tmp_path / "sample.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def engine():
    return MaskingEngine(deny_lists=DENY, enable_ner=False)


@pytest.fixture
def vault(vaults_dir):
    return Vault.create("sample.docx", vaults_dir=vaults_dir)


class TestExtraction:
    def test_all_text_captured(self, fixture_docx):
        parse = parse_docx(fixture_docx)
        combined = parse.spanned.text
        for fragment in ["Acme Corp Review", "John Smith", "Acme ", "Corp expands",
                         "Bidco Ltd", "Closing line"]:
            assert fragment in combined

    def test_no_cross_paragraph_bleed(self, fixture_docx):
        # paragraphs must be separated so entities can't falsely match across them
        parse = parse_docx(fixture_docx)
        assert "board.Our client" not in parse.spanned.text


class TestMaskedDocx:
    def test_split_run_entity_detected(self, fixture_docx, engine, vault, tmp_path):
        parse = parse_docx(fixture_docx)
        new_texts = engine.mask_spanned(parse.spanned, vault)
        out = tmp_path / "masked.docx"
        write_masked_docx(parse, new_texts, out)
        text = _doc_text(out)
        assert "Acme" not in text  # the split "Acme " + "Corp" was caught
        assert vault.resolve("⟦ORG_001⟧") in ("Acme Corp", "Bidco Ltd")

    def test_zero_leakage_in_masked_docx(self, fixture_docx, engine, vault, tmp_path):
        parse = parse_docx(fixture_docx)
        new_texts = engine.mask_spanned(parse.spanned, vault)
        out = tmp_path / "masked.docx"
        write_masked_docx(parse, new_texts, out)
        text = _doc_text(out)
        for term in ALL_TERMS:
            assert term not in text, f"leaked: {term}"
        assert TOKEN_PATTERN.search(text)

    def test_formatting_preserved(self, fixture_docx, engine, vault, tmp_path):
        parse = parse_docx(fixture_docx)
        new_texts = engine.mask_spanned(parse.spanned, vault)
        out = tmp_path / "masked.docx"
        write_masked_docx(parse, new_texts, out)
        reopened = docx.Document(str(out))
        # the bold "John Smith" run now carries a PERSON token and is still bold
        bold_runs = [
            r for p in reopened.paragraphs for r in p.runs if r.bold
        ]
        assert any(TOKEN_PATTERN.search(r.text) for r in bold_runs)

    def test_table_cells_masked(self, fixture_docx, engine, vault, tmp_path):
        parse = parse_docx(fixture_docx)
        new_texts = engine.mask_spanned(parse.spanned, vault)
        out = tmp_path / "masked.docx"
        write_masked_docx(parse, new_texts, out)
        reopened = docx.Document(str(out))
        table_text = "\n".join(
            cell.text for row in reopened.tables[0].rows for cell in row.cells
        )
        assert "Bidco Ltd" not in table_text
        assert "Company" in table_text  # non-sensitive header untouched


class TestMarkdownRendition:
    def test_masked_markdown_from_docx(self, fixture_docx, engine, vault):
        masked_md, warnings = mask_docx_to_markdown(fixture_docx, engine, vault)
        for term in ALL_TERMS:
            assert term not in masked_md, f"leaked: {term}"
        assert masked_md.lstrip().startswith("# ")  # heading preserved
        assert "|" in masked_md  # table rendered
        assert warnings == []


class TestRestore:
    def test_restore_docx_round_trip(self, fixture_docx, engine, vault, tmp_path):
        original_text = _doc_text(fixture_docx)
        parse = parse_docx(fixture_docx)
        new_texts = engine.mask_spanned(parse.spanned, vault)
        masked_path = tmp_path / "masked.docx"
        write_masked_docx(parse, new_texts, masked_path)

        restored_path = tmp_path / "restored.docx"
        unresolved = restore_docx(masked_path, vault, restored_path)
        assert unresolved == []
        assert _doc_text(restored_path) == original_text
